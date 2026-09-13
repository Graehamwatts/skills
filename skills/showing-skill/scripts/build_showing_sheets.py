"""
Build the client + working showing-tour spreadsheets from a JSON spec.

Usage:
    python build_showing_sheets.py <spec.json> <output_dir>

See examples/spec-example.json for the input schema. This script only does the
mechanical parts (time-window math, spreadsheet formatting). Reading the MLS
PDFs/screenshots, deciding which properties need agent contact, drafting the
emails, and identifying which contacts need a TEXT instead of an email are all
judgment calls Claude makes when building the spec -- see SKILL.md.
"""
import sys
import json
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY, LIGHT_GRAY, WHITE, RED, GREEN = "1A365D", "F2F2F2", "FFFFFF", "C00000", "1E7B34"
FONT_NAME = "Arial"


def round5(dt):
    discard = timedelta(minutes=dt.minute % 5, seconds=dt.second, microseconds=dt.microsecond)
    dt -= discard
    if discard >= timedelta(minutes=2.5):
        dt += timedelta(minutes=5)
    return dt


def fmt_time(dt, with_ampm=True):
    # dt.strftime("%-I") isn't portable to Windows Python builds; strip the
    # leading zero manually instead.
    hour = dt.strftime("%I").lstrip("0") or "12"
    out = f'{hour}:{dt.strftime("%M")}'
    return f'{out} {dt.strftime("%p")}' if with_ampm else out


def compute_windows(spec):
    """Chain arrival windows forward: window = [nominal, nominal + window_minutes].
    Stop 1's nominal start is the user-given start_time (not re-derived).
    Later stops chain off the PREVIOUS stop's rounded nominal start, not its
    window end -- this matches the "give a 30-min range" convention without
    the range width compounding across stops."""
    show_min = spec.get("showing_minutes", 15)
    win_min = spec.get("window_minutes", 30)
    base_date = spec["tour_date"]
    nominal = datetime.strptime(f'{base_date} {spec["start_time"]}', "%Y-%m-%d %H:%M")
    for i, p in enumerate(spec["properties"]):
        if i == 0:
            start = nominal
        else:
            prev_drive = spec["properties"][i - 1].get("drive_to_next_min", 0)
            start = round5(nominal + timedelta(minutes=show_min + prev_drive))
            nominal = start
        end = start + timedelta(minutes=win_min)
        same_period = start.strftime("%p") == end.strftime("%p")
        p["_window"] = f'{fmt_time(start, with_ampm=not same_period)} - {fmt_time(end)}'
    return spec


def money(v):
    return f"${v:,.0f}"


def style_header_row(ws, row, ncols, height=34):
    ws.row_dimensions[row].height = height
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=NAVY))


def style_title_block(ws, ncols, title, subtitle):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    t = ws.cell(row=1, column=1, value=title)
    t.font = Font(name=FONT_NAME, bold=True, size=16, color=NAVY)
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font = Font(name=FONT_NAME, italic=True, size=10, color="595959")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18


def apply_borders_and_banding(ws, first_row, last_row, ncols):
    thin = Side(style="thin", color="D9D9D9")
    for r in range(first_row, last_row + 1):
        fill = PatternFill("solid", fgColor=LIGHT_GRAY) if (r - first_row) % 2 == 1 else PatternFill("solid", fgColor=WHITE)
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if not cell.fill.fgColor.rgb or cell.fill.fgColor.rgb == "00000000":
                cell.fill = fill


def build_client_sheet(spec, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "Showing Schedule"
    headers = ["Stop", "Address", "City / Zip", "Time Window", "Status",
               "List Price", "Beds", "Baths", "SqFt", "Lot SqFt", "DOM", "Drive to Next", "Notes"]
    ncols = len(headers)
    total_mi = sum(p.get("drive_to_next_mi", 0) for p in spec["properties"])
    total_min = sum(p.get("drive_to_next_min", 0) for p in spec["properties"])
    style_title_block(ws, ncols,
                       f'Property Tour for {spec["client_name"]} — {spec["tour_date_display"]}',
                       f'Start {spec["properties"][0]["_window"].split(" - ")[0]} at {spec["properties"][0]["address"]}  |  '
                       f'{len(spec["properties"])} homes, ~{total_mi:.1f} mi / {total_min} min total drive time  |  '
                       f'Prepared by {spec["agent"]["name"]}, {spec["agent"]["brokerage"]}')
    header_row = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=h)
    style_header_row(ws, header_row, ncols)
    r = header_row + 1
    for i, p in enumerate(spec["properties"], start=1):
        drive = f'{p["drive_to_next_mi"]} mi / {p["drive_to_next_min"]} min' if p.get("drive_to_next_min") else "—  (last stop)"
        vals = [i, p["address"], f'{p["city"]}, {p["zip"]}', p["_window"], p["occupancy"],
                money(p["price"]), p["beds"], p["baths"], p["sqft"], p["lot"], p["dom"], drive, p["client_notes"]]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
        r += 1
    last_row = r - 1
    apply_borders_and_banding(ws, header_row + 1, last_row, ncols)
    widths = [6, 20, 20, 15, 14, 12, 6, 6, 8, 9, 6, 15, 46]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in (1, 4, 5, 7, 8, 9, 10, 11):
        for row in ws.iter_rows(min_row=header_row + 1, max_row=last_row, min_col=col, max_col=col):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    footer_row = last_row + 2
    f = ws.cell(row=footer_row, column=1,
                value="Times are approximate arrival windows, not exact — running early or late by a few minutes is normal.")
    f.font = Font(name=FONT_NAME, italic=True, size=9, color="7F7F7F")
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=ncols)
    path = f'{out_dir}/{spec["client_name"]} Showing Schedule - CLIENT.xlsx'
    wb.save(path)
    return path


def build_working_sheet(spec, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "Working Sheet"
    headers = ["Stop", "Address", "City / Zip", "Time Window", "Occupancy",
               "List Price", "Beds", "Baths", "SqFt", "Lot SqFt", "DOM", "MLS #", "Drive to Next",
               "Showing / Access Instructions", "Listing Agent", "LA Phone", "LA Email", "Co-Listing Agent",
               "Contact Status", "Flag"]
    ncols = len(headers)
    total_mi = sum(p.get("drive_to_next_mi", 0) for p in spec["properties"])
    total_min = sum(p.get("drive_to_next_min", 0) for p in spec["properties"])
    style_title_block(ws, ncols,
                       f'{spec["client_name"].upper()} SHOWINGS — WORKING SHEET (agent info) — {spec["tour_date_display"]}',
                       f'Start {spec["properties"][0]["_window"].split(" - ")[0]} at {spec["properties"][0]["address"]}  |  '
                       f'{len(spec["properties"])} homes, ~{total_mi:.1f} mi / {total_min} min total drive time  |  Internal use only — do not print for client')
    header_row = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=h)
    style_header_row(ws, header_row, ncols, height=40)
    r = header_row + 1
    for i, p in enumerate(spec["properties"], start=1):
        drive = f'{p["drive_to_next_mi"]} mi / {p["drive_to_next_min"]} min' if p.get("drive_to_next_min") else "—  (last stop)"
        la = p["listing_agent"]
        co = p.get("co_listing_agent")
        co_str = f'{co["name"]} (co-LA) — {co["phone"]} — {co["email"]}' if co else ""
        vals = [i, p["address"], f'{p["city"]}, {p["zip"]}', p["_window"], p["occupancy"],
                money(p["price"]), p["beds"], p["baths"], p["sqft"], p["lot"], p["dom"], p["mls"], drive,
                p["access_instructions"], f'{la["name"]}, {la["brokerage"]}', la["phone"], la["email"], co_str,
                p["contact_status"], p.get("flag", "")]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
        if p.get("flag"):
            ws.cell(row=r, column=20).font = Font(name=FONT_NAME, size=10, bold=True, color=RED)
        occ_cell = ws.cell(row=r, column=5)
        occ_cell.font = Font(name=FONT_NAME, size=10, bold=True, color=RED if p["occupancy"] == "Owner Occupied" else GREEN)
        r += 1
    last_row = r - 1
    apply_borders_and_banding(ws, header_row + 1, last_row, ncols)
    widths = [6, 20, 20, 15, 14, 11, 6, 6, 7, 9, 6, 13, 13, 40, 22, 14, 24, 30, 34, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    texts = [p for p in spec["properties"] if p.get("contact_channel") in ("text", "call_or_text")]
    if texts:
        ws3 = wb.create_sheet("Texts To Send")
        ws3["A1"] = "TEXTS TO SEND (Claude/email can't cover these — a human needs to send them)"
        ws3["A1"].font = Font(name=FONT_NAME, bold=True, size=12, color=RED)
        ws3.merge_cells("A1:C1")
        for i, h in enumerate(["Property / Time", "Send To", "Suggested Message"], start=1):
            ws3.cell(row=3, column=i, value=h)
        style_header_row(ws3, 3, 3)
        rr = 4
        for p in texts:
            la = p["listing_agent"]
            msg = p.get("suggested_text") or (
                f'Hi {la["name"].split()[0]}, this is {spec["agent"]["name"]} - showing {p["address"]} today '
                f'between {p["_window"]}, could you help with access? Thanks!')
            ws3.cell(row=rr, column=1, value=f'{p["address"]} — {p["_window"]}')
            ws3.cell(row=rr, column=2, value=f'{la["name"]}: {la["phone"]}')
            ws3.cell(row=rr, column=3, value=msg)
            rr += 1
        apply_borders_and_banding(ws3, 4, rr - 1, 3)
        ws3.column_dimensions["A"].width = 30
        ws3.column_dimensions["B"].width = 26
        ws3.column_dimensions["C"].width = 70

    path = f'{out_dir}/{spec["client_name"]} Showing Schedule - WORKING (agent info).xlsx'
    wb.save(path)
    return path


def main():
    spec_path, out_dir = sys.argv[1], sys.argv[2]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    spec.setdefault("tour_date_display", spec["tour_date"])
    compute_windows(spec)
    c = build_client_sheet(spec, out_dir)
    a = build_working_sheet(spec, out_dir)
    print("Saved:", c)
    print("Saved:", a)


if __name__ == "__main__":
    main()
