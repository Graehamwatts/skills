# -*- coding: utf-8 -*-
"""intake_sheet: one "<address> - View Activity" workbook -> raw.json.

    python intake_sheet.py <export.xlsx> --folder <reports>\\seller-reports\\<slug> --as-of 2026-10-12

The workbook is the xlsx export of the sheet the client care coordinator keeps by hand. This script reads
every tab, every row and every dated note exactly as typed and writes raw.json (see
references/data-contract.md, section 3). It interprets nothing beyond dates: who is who, what a note means
and how it is worded for a seller are decided later, in curation.json.

What it is built to survive:
  - tabs are found by keyword (Glide or disclosure, showing or supra, open house, inquiry, offers), never
    by exact name. A missing or unrecognised tab is reported loudly, not skipped in silence.
  - the header row is found wherever it sits, and columns are mapped from the header text.
  - every Notes column is read, out to the last used column, including unnamed spill columns.
  - dates typed without a year are placed by the listing's life, never by today's year.
  - no row is dropped silently: every row that is not read as a row is listed in skipped_rows with a reason.
  - rows that may be the same person are proposed as merge CANDIDATES for Claude to settle. Nothing is merged here.

raw.json is PRIVATE (full names, phone numbers, notes as typed). It is refused inside a git repo.
Needs openpyxl. No real client data belongs in this file.
"""
import argparse
import datetime
import os
import re
import site
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)
if site.getusersitepackages() not in sys.path:
    sys.path.append(site.getusersitepackages())

import sr_contract as C          # noqa: E402
import sr_files as F             # noqa: E402
import sr_sheet_text as T        # noqa: E402

# header text -> field. The first rule that matches wins, so the narrow rules come first.
HEADER_RULES = (
    ("agent_no", r"^(agent|agt)\s*(#|no\.?|number)$"),
    ("buyer_no", r"^buyer\s*(#|no\.?|number)$"),
    ("email_addr", r"^(agent\s*)?e-?mail\s*(address|add\.?|addr\.?)$|^agent e-?mail$"),
    ("phone", r"^(phone|cell|mobile|contact)(\s*(number|no\.?|#))?$|^phone\b.*"),
    ("buyer", r"^buyer'?s?\s*name$|^buyers?(/s)?$|^client\s*name$"),
    ("name", r"^(agent|showing agent|buyer'?s? agent|host|hosting agent|agent to host)(\s*name)?$|^name$|^agent\s*/\s*buyer(\s*name)?$"),
    ("next_followup_date", r"next\s*(follow|f/?u|touch|call|step)|follow[\s\-]*up\s*(date|due|on|by)|f/?u\s*(date|due)"),
    ("price_reaction", r"price\s*(reaction|feedback|comments?|response|view|opinion|discussed|came up)|reaction\s*to\s*price|^price\??$"),
    ("interest", r"interest"),
    ("offer_intent", r"offer\s*(intent|intention|likel\w*|plan\w*|\?)|^intent$|will\s*(they\s*)?(write|offer)"),
    ("visitor_count", r"^(#|no\.?|number|count|total)\s*(of\s*)?(visitors?|groups?|guests?|parties|people|attendees)"
                      r"|(visitors?|groups?|guests?|attend\w*|traffic)\s*(/\s*\w+\s*)?(count|#|number|total)"
                      r"|^(visitors?|groups?|guests?|attendance|traffic|count|head\s*count)$"),
    ("when", r"^date\s*(&|and|/)\s*time$|^date of\b.*|^(showing|open house|oh|tour|viewing|received)\s*date.*$|^date(\s*received)?$"),
    ("feedback", r"^(agent\s*|buyer\s*|host\s*)?(feedback|comments?|remarks?)$|^offer summary$|^summary$"),
    ("last_email", r"^(last\s*)?e-?mail(ed)?$"),
    ("last_call", r"^(last\s*)?call(ed)?$|^phone call$"),
    ("last_text", r"^(last\s*)?(text|sms)(ed)?$"),
    ("notes", r"^(follow[\s\-]*up\s*)?notes?$|^follow[\s\-]*ups?$|^follow[\s\-]*up log$"),
)
_RULES = tuple((k, re.compile(rx, re.I)) for k, rx in HEADER_RULES)
MULTI = ("feedback", "notes")                     # fields that may span several columns
EXPECTED_TABS = ("disclosure", "showing", "open_house")
TAG_FIELDS = ("price_reaction", "interest", "offer_intent", "next_followup_date")
_HOST_HINT = re.compile(r"open\s*house\s*(host|house)|hosted?\s+(the\s+|an?\s+)?open\s*house|helped\s+.*open\s*house|\(host\)", re.I)
_LEAD_HINT = re.compile(r"buyer\s*lead|\(buyer\)|is a buyer|no agent|unrepresented", re.I)
_XREF = re.compile(r"^refer to (the )?showing activity\b", re.I)


def col_letter(i):
    """0 -> 'A', 26 -> 'AA'."""
    s, i = "", i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def header_field(text):
    if not isinstance(text, str):
        return None
    t = " ".join(text.lower().replace("\n", " ").split()).strip(" :*")
    for key, rx in _RULES:
        if rx.search(t):
            return key
    return None


def classify_tab(title, fields, header_texts):
    """The tab kind, from keywords in its title and then from its headers. Never from an exact name."""
    t = " ".join(str(title or "").lower().split())
    heads = " | ".join(h.lower() for h in header_texts if isinstance(h, str))
    if "open house" in t or "openhouse" in t or re.search(r"\boh\b", t) or "date of open house" in heads:
        return "open_house"
    if "glide" in t or "disclosure" in t:
        return "disclosure"
    if "inquir" in t or "enquir" in t:
        return "inquiry"
    if "offer" in t:
        return "offers"
    if "showing" in t or "supra" in t or "tour" in t or "viewing" in t or re.search(r"date of (viewing|showing)", heads):
        return "showing"
    return "other"


def find_header(grid):
    """(row index, {column index: field}) for the first row in the top 15 that reads as a header."""
    for i, row in enumerate(grid[:15]):
        fields = {}
        for ci, v in enumerate(row):
            k = header_field(v)
            if k and (k in MULTI or k not in fields.values()):
                fields[ci] = k
        got = set(fields.values())
        if len(got) >= 2 and got & set(("name", "feedback", "when", "buyer")):
            return i, fields
    return None, {}


def count_unit(header):
    h = (header or "").lower()
    hits = [u for u, rx in (("groups", "group"), ("visitors", "visitor"), ("guests", "guest"), ("parties", "part(y|ies)"),
                            ("people", "people|attend|head")) if re.search(rx, h)]
    return hits[0] if len(hits) == 1 else None


class Intake(object):
    """Reads one workbook. Keeps the running lists every tab adds to: warnings, skipped rows, taken keys."""

    def __init__(self, ctx):
        self.ctx = ctx
        self.warnings, self.skipped, self.taken = [], [], {}

    def warn(self, text):
        if text not in self.warnings:
            self.warnings.append(text)

    def unique(self, key, taken=None):
        taken = self.taken if taken is None else taken
        taken[key] = taken.get(key, 0) + 1
        return key if taken[key] == 1 else "%s~%d" % (key, taken[key])

    # -------------------------------------------------------------------------------------------- one tab
    def read_tab(self, ws, source_id):
        grid = [[T.clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
        while grid and not any(v is not None for v in grid[-1]):
            grid.pop()
        width = max([max([i for i, v in enumerate(r) if v is not None] or [-1]) for r in grid] or [-1]) + 1
        hdr_i, fields = find_header(grid)
        merged = [str(r) for r in ws.merged_cells.ranges]
        hidden = sorted(set(col_letter(c - 1) for d in ws.column_dimensions.values() if d.hidden
                            for c in range(d.min or 1, (d.max or d.min or 1) + 1) if c <= width))
        tab = {"source_id": source_id, "title": ws.title, "kind": "other", "state": ws.sheet_state,
               "header_row": None, "headers": {}, "columns_found": {}, "columns_unrecognised": {},
               "hidden_columns": hidden, "merged": merged, "rows": []}
        where = "tab %r" % ws.title
        if hdr_i is None:
            tab["kind"] = classify_tab(ws.title, {}, [])
            if grid:
                self.warn("%s: no header row found, so none of its %d rows could be read" % (where, len(grid)))
            for ri, row in enumerate(grid):
                if any(v is not None for v in row):
                    self.skip(ws.title, ri + 1, "no header row found in this tab", row)
            return tab
        header = list(grid[hdr_i]) + [None] * (width - len(grid[hdr_i]))
        for rng in ws.merged_cells.ranges:                       # a header merged across columns names all of them
            if rng.min_row <= hdr_i + 1 <= rng.max_row:
                top = header[rng.min_col - 1] if rng.min_col - 1 < len(header) else None
                for c in range(rng.min_col, rng.max_col + 1):
                    if c - 1 < len(header) and header[c - 1] is None and top is not None:
                        header[c - 1] = top
                        if header_field(top) in MULTI:
                            fields[c - 1] = header_field(top)
        named = [i for i, h in enumerate(header) if h is not None]
        last_named = max(named) if named else -1
        tab["header_row"] = hdr_i + 1
        tab["headers"] = dict((col_letter(i), T.cell_json(header[i])) for i in named)
        tab["kind"] = kind = classify_tab(ws.title, fields, [h for h in header if isinstance(h, str)])
        cols = {}
        for ci, k in sorted(fields.items()):
            cols.setdefault(k, []).append(ci)
        used = set(i for r in grid[hdr_i + 1:] for i, v in enumerate(r) if v is not None)
        for c in list(cols.get("notes", [])):                    # unnamed columns right after Notes are more notes,
            j = c + 1                                            # even when another header follows further right
            while j <= last_named and header[j] is None:
                if j in used:
                    cols["notes"].append(j)
                    fields[j] = "notes"
                j += 1
        if "notes" in cols:
            cols["notes"] = sorted(set(cols["notes"]))
        spill = sorted(i for i in used if i > last_named)
        if spill:
            cols.setdefault("notes", [])
            if not [c for c in cols["notes"] if c <= last_named]:
                self.warn("%s: text in unnamed columns %s to %s is read as notes (the tab has no Notes header)"
                          % (where, col_letter(spill[0]), col_letter(spill[-1])))
            cols["notes"] = sorted(set(cols["notes"]) | set(spill))
        for i in sorted(used):
            if i <= last_named and i not in fields:
                name = header[i] if isinstance(header[i], str) else None
                tab["columns_unrecognised"][col_letter(i)] = name
                self.warn("%s: column %s (%s) is not a column this parser knows; its cells are kept, not interpreted"
                          % (where, col_letter(i), repr(name) if name else "no header"))
        tab["columns_found"] = dict((k, [col_letter(c) for c in v] if k in MULTI else col_letter(v[0])) for k, v in cols.items())
        if kind == "other":
            self.warn("%s is not a tab this parser recognises (expected a Glide or disclosure, showing, open house, "
                      "inquiry or offers tab). Its rows are kept as 'other' and are in no count" % where)
        if ws.sheet_state != "visible":
            self.warn("%s is a hidden tab. It was read like the others" % where)
        if "name" not in cols and "buyer" not in cols:
            self.warn("%s: no name column was found" % where)
        for ri in range(hdr_i):
            if any(v is not None for v in grid[ri]):
                self.skip(ws.title, ri + 1, "above the header row (a title or a note)", grid[ri])
        unit = count_unit(header[cols["visitor_count"][0]]) if "visitor_count" in cols else None
        for ri in range(hdr_i + 1, len(grid)):
            row = list(grid[ri]) + [None] * (width - len(grid[ri]))
            if not any(v is not None for v in row):
                continue
            rec = self.read_row(tab, cols, row, ri + 1, unit)
            if rec is not None:
                tab["rows"].append(rec)
        return tab

    def skip(self, tab_title, row_number, reason, row):
        cells = dict((col_letter(i), T.cell_json(v)) for i, v in enumerate(row) if v is not None)
        self.skipped.append({"tab": tab_title, "row": row_number, "reason": reason, "cells": cells})

    # -------------------------------------------------------------------------------------------- one row
    def read_row(self, tab, cols, row, number, unit):
        ctx, kind, where = self.ctx, tab["kind"], "tab %r row %d" % (tab["title"], number)

        def cell(field):
            return row[cols[field][0]] if field in cols else None

        filled = set(i for i, v in enumerate(row) if v is not None)
        numbering = set(cols.get("agent_no", []) + cols.get("buyer_no", []))
        if filled <= numbering:
            self.skip(tab["title"], number, "only a row number: no name, contact detail or note", row)
            return None
        heads = [h for h in tab["headers"].values() if isinstance(h, str)]
        texts = [v for v in row if isinstance(v, str)]
        if len(texts) >= 2 and all(t in heads for t in texts):
            self.skip(tab["title"], number, "repeats the header row", row)
            return None
        name_raw = cell("name")
        name_raw = str(name_raw) if name_raw is not None and not isinstance(name_raw, (datetime.date, datetime.datetime)) else None
        buyer_raw = cell("buyer")
        buyer_raw = str(buyer_raw) if buyer_raw is not None and not isinstance(buyer_raw, (datetime.date, datetime.datetime)) else None
        invited = C.invited_by(name_raw) if name_raw else None
        contact_cells = [cell("phone"), cell("email_addr"), name_raw]
        phones, emails = T.phones_of(*contact_cells), T.emails_of(*[c for c in contact_cells if isinstance(c, str)])
        when_v = cell("when")
        rec = {"key": None, "row": number,
               "cells": dict((col_letter(i), T.cell_json(v)) for i, v in enumerate(row) if v is not None),
               "name_raw": name_raw, "buyer_raw": buyer_raw, "invited_by_raw": invited,
               "agent_no": T.cell_json(cell("agent_no")), "buyer_no": T.cell_json(cell("buyer_no")),
               "phones": phones, "emails": emails,
               "when_raw": T.cell_json(when_v), "when": [], "last_touch": {}, "last_touch_all": {},
               "sheet_tags": dict((k, None) for k in TAG_FIELDS), "sheet_tags_raw": {},
               "visitor_count": None, "hints": [], "entries": []}
        if kind != "open_house":
            rec["when"] = [dict((k, w[k]) for k in ("date", "time_text", "date_quality", "weekday", "year_rule", "raw"))
                           for w in T.parse_when(when_v, ctx)]
            if when_v is not None and not any(w["date"] for w in rec["when"]):
                self.warn("%s: the date cell holds no date that could be read" % where)
        for ch in ("email", "call", "text"):
            dates, problem = T.parse_last_touch(cell("last_" + ch), ctx)
            rec["last_touch"][ch] = dates[-1] if dates else None
            if len(dates) > 1:
                rec["last_touch_all"][ch] = dates
            if problem:
                self.warn("%s: the %s column: %s" % (where, ch.title(), problem))
        for k in TAG_FIELDS:
            v = cell(k)
            if v is not None:
                rec["sheet_tags_raw"][k] = T.cell_json(v)
                rec["sheet_tags"][k], problem = T.read_sheet_tag(k, v, ctx)
                if problem:
                    self.warn("%s: %s. Left for Claude to read from the note" % (where, problem))
        extra = ""
        if kind == "open_house":
            extra = str(rec["when_raw"] or "")
        elif invited and rec["buyer_no"] is not None:
            extra = "buyer %s" % rec["buyer_no"]
        elif not (name_raw or buyer_raw or phones or emails):
            extra = " ".join(str(v) for v in row if isinstance(v, str))[:60]
        rec["key"] = self.unique(C.row_key(kind, name_raw or buyer_raw or "", phones, emails, extra))
        seen = {}
        for field, plural in (("feedback", "feedback"), ("note", "notes")):
            if kind == "open_house" and field == "feedback":
                continue                                           # an open house recap is read into sessions below
            for ci in cols.get(plural, []):
                v = row[ci]
                if v is None:
                    continue
                if not isinstance(v, str):
                    v = str(T.cell_json(v))
                for e in T.split_dated(v, ctx):
                    e.update(field=field, col=col_letter(ci))
                    e["key"] = self.unique(C.entry_key(rec["key"], field, e["date"], e["text"]), seen)
                    if e["year_rule"] == "outside_listing_life":
                        self.warn("%s: the note dated %s falls outside the listing's life and was placed on %s"
                                  % (where, e["date_raw"], e["date"]))
                    if e.get("mid_line_dates"):
                        self.warn("%s: a second date (%s) sits in the middle of a line. It was left inside the note dated %s"
                                  % (where, ", ".join(e["mid_line_dates"]), e["date_raw"] or "undated"))
                    rec["entries"].append(dict((k, e[k]) for k in ("key", "field", "col", "date", "date_raw", "year_rule",
                                                                     "text", "reply_lines", "mid_line_dates") if k in e))
        if kind == "open_house":
            recap = "\n".join(str(row[ci]) for ci in cols.get("feedback", []) if row[ci] is not None) or None
            count = cell("visitor_count")
            rec["visitor_count"] = T.cell_json(count)
            sessions, problems = T.parse_sessions(when_v, recap, ctx, count, unit)
            for s in sessions:
                s["key"] = "%s#%d" % (rec["key"], s["index"])
            rec["sessions"], rec["recap_raw"], rec["visitor_count_unit"] = sessions, recap, unit
            for p in problems:
                self.warn("%s: %s" % (where, p))
            if any(s["date"] is None for s in sessions):
                self.warn("%s: an open house session has no date typed. It needs a date from a person before it can be counted" % where)
        all_text = " ".join(e["text"] for e in rec["entries"]) + " " + (rec.get("recap_raw") or "") + " " + (name_raw or "")
        if kind != "open_house" and _HOST_HINT.search(all_text):
            rec["hints"].append("open_house_host")
        if _LEAD_HINT.search(all_text):
            rec["hints"].append("buyer_lead")
        if invited:
            rec["hints"].append("invited_by")
        if not name_raw and not buyer_raw:
            rec["hints"].append("no_name")
        if any(_XREF.match(e["text"]) for e in rec["entries"]):
            rec["hints"].append("refers_to_showing_tab")
        return rec


# ---------------------------------------------------------------------------------------------
# merge candidates: proposals for a reviewer, never applied here
# ---------------------------------------------------------------------------------------------
def merge_candidates(tabs):
    """Rows that may be the same person. Phone digits link first, then the name key. Returns
    (candidates, inviter_candidates): candidates is a list of {rows, basis, tabs}; inviter_candidates maps an
    'Invited by' row to the rows whose name matches the inviter."""
    rows = [(t, r) for t in tabs for r in t["rows"] if t["kind"] not in ("offers", "other")]
    people = [(t, r) for t, r in rows if not r.get("invited_by_raw")]
    parent, why = {}, {}

    def find(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b, basis):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb
        why.setdefault(frozenset((a, b)), basis)

    by_phone, by_name = {}, {}
    for _, r in people:
        find(r["key"])
        for p in r["phones"]:
            if p in by_phone:
                union(r["key"], by_phone[p], "phone")
            by_phone[p] = r["key"]
    for _, r in people:
        for nm in (r["name_raw"], r["buyer_raw"]):
            k = T.name_key(nm)
            if k:
                if k in by_name and find(by_name[k]) != find(r["key"]):
                    union(r["key"], by_name[k], "name")
                by_name.setdefault(k, r["key"])
    groups = {}
    for t, r in people:
        groups.setdefault(find(r["key"]), []).append((t, r))
    out = []
    for members in groups.values():
        if len(members) < 2:
            continue
        keys = [r["key"] for _, r in members]
        bases = sorted(set(b for pair, b in why.items() if pair <= set(keys)))
        out.append({"rows": keys, "basis": "name" if "name" in bases else "phone",
                    "tabs": sorted(set(t["kind"] for t, _ in members)),
                    "names_differ": len(set(T.name_key(r["name_raw"] or r["buyer_raw"]) for _, r in members) - set([None])) > 1})
    out.sort(key=lambda g: g["rows"])
    inviters = {}
    for t, r in rows:
        if r.get("invited_by_raw"):
            k = T.name_key(r["invited_by_raw"])
            inviters[r["key"]] = sorted(set(x["key"] for _, x in people if k and T.name_key(x["name_raw"]) == k))
    return out, inviters


# ---------------------------------------------------------------------------------------------
# the workbook
# ---------------------------------------------------------------------------------------------
def parse_workbook(path, list_date, as_of, slug, source=None):
    """Reads the workbook and returns the raw.json dict."""
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl is needed to read the workbook: python -m pip install --user openpyxl")
    ctx = T.DateCtx(list_date, as_of)
    wb = openpyxl.load_workbook(path, data_only=True)
    job = Intake(ctx)
    src = {"source_id": "s1", "kind": "activity_sheet_live", "title": None, "drive_id": None,
           "file_name": os.path.basename(path), "modified_utc": None,
           "exported_at": datetime.datetime.fromtimestamp(os.path.getmtime(path)).replace(microsecond=0).isoformat(),
           "as_of": C.iso(as_of), "sha256": F.sha256_file(path)}
    src.update(source or {})
    if src["kind"] not in C.SOURCE_KINDS:
        raise SystemExit("unknown source kind %r" % src["kind"])
    tabs = [job.read_tab(ws, src["source_id"]) for ws in wb.worksheets]
    found = {}
    for t in tabs:
        found.setdefault(t["kind"], []).append(t["title"])
    missing = [k for k in EXPECTED_TABS if k not in found]
    for k in missing:
        job.warn("no %s tab was found in this workbook (tabs: %s)" % (k.replace("_", " "), ", ".join(repr(t["title"]) for t in tabs) or "none"))
    cands, inviters = merge_candidates(tabs)
    rows = [r for t in tabs for r in t["rows"]]
    entries = [e for r in rows for e in r["entries"]]
    sessions = [s for r in rows for s in r.get("sessions") or []]
    stats = {
        "tabs": len(tabs), "rows_total": len(rows),
        "rows_by_tab_kind": dict((k, sum(len(t["rows"]) for t in tabs if t["kind"] == k)) for k in sorted(found)),
        "rows_skipped": len(job.skipped), "invited_by_rows": len([r for r in rows if r.get("invited_by_raw")]),
        "rows_without_name": len([r for r in rows if "no_name" in r["hints"]]),
        "entries_total": len(entries), "entries_undated": len([e for e in entries if not e["date"]]),
        "feedback_entries": len([e for e in entries if e["field"] == "feedback"]),
        "note_entries": len([e for e in entries if e["field"] == "note"]),
        "dates_outside_listing_life": len([e for e in entries if e.get("year_rule") == "outside_listing_life"]),
        "rows_with_sheet_tags": len([r for r in rows if any(v is not None for v in r["sheet_tags"].values())]),
        "open_house_sessions": len(sessions),
        "open_house_sessions_undated": len([s for s in sessions if not s["date"]]),
        "open_house_sessions_with_count": len([s for s in sessions if s["groups"] is not None]),
        "merge_candidates": len(cands), "warnings": len(job.warnings),
    }
    return {
        "schema_version": C.SCHEMA_VERSION,
        "WARNING": "PRIVATE. Full names, phone numbers, emails and notes as typed. Never commit or share.",
        "slug": slug, "parsed_at": datetime.datetime.now().astimezone().replace(microsecond=0).isoformat(),
        "list_date": C.iso(list_date), "as_of": C.iso(as_of),
        "sources": [src], "tabs": tabs, "tabs_found": found, "tabs_missing": missing,
        "merge_candidates": cands, "inviter_candidates": inviters, "skipped_rows": job.skipped,
        "stats": stats, "warnings": job.warnings,
    }


def changes_since(old, new):
    """What moved between two parses, by key: counts only."""
    def keys(raw):
        rows = set(r["key"] for t in raw.get("tabs") or [] for r in t["rows"])
        ents = set(e["key"] for t in raw.get("tabs") or [] for r in t["rows"] for e in r["entries"])
        return rows, ents
    (r0, e0), (r1, e1) = keys(old), keys(new)
    return {"previous_parse": old.get("parsed_at"), "rows_new": len(r1 - r0), "rows_gone": len(r0 - r1),
            "entries_new_or_edited": len(e1 - e0), "entries_gone_or_edited": len(e0 - e1)}


def summary(raw):
    """A plain summary for the operator. Counts only: no names, phones or notes."""
    lines, st = [], raw["stats"]
    for t in raw["tabs"]:
        lines.append("  tab %-34r %-11s header row %s, %d rows, columns: %s"
                     % (t["title"], t["kind"], t["header_row"], len(t["rows"]),
                        ", ".join("%s=%s" % (k, "+".join(v) if isinstance(v, list) else v) for k, v in sorted(t["columns_found"].items()))))
    lines.append("  rows %d, dated notes %d, undated notes %d, open house sessions %d (%d with a count, %d undated)"
                 % (st["rows_total"], st["entries_total"] - st["entries_undated"], st["entries_undated"],
                    st["open_house_sessions"], st["open_house_sessions_with_count"], st["open_house_sessions_undated"]))
    lines.append("  merge candidates: %d. Rows not read as rows: %d (each listed in skipped_rows with its reason)"
                 % (st["merge_candidates"], st["rows_skipped"]))
    if raw.get("since_last_parse"):
        c = raw["since_last_parse"]
        lines.append("  since the last parse: %d new rows, %d rows gone, %d new or edited notes"
                     % (c["rows_new"], c["rows_gone"], c["entries_new_or_edited"]))
    return "\n".join(lines)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Parse one View Activity workbook into raw.json (PRIVATE output).")
    ap.add_argument("xlsx", help="the xlsx export of the activity sheet")
    ap.add_argument("--folder", required=True, help="the listing folder: <reports>\\seller-reports\\<slug>")
    ap.add_argument("--list-date", help="YYYY-MM-DD. Default: list_date in the folder's listing.json")
    ap.add_argument("--as-of", help="YYYY-MM-DD the sheet was exported or last edited. Default: today")
    ap.add_argument("--slug", help="default: slug in listing.json, else the folder name")
    ap.add_argument("--kind", default="activity_sheet_live", choices=("activity_sheet_live", "activity_sheet_file"),
                    help="activity_sheet_live for a fresh export of the live sheet, activity_sheet_file for a copy someone attached")
    ap.add_argument("--title", help="the sheet's title, for the record")
    ap.add_argument("--drive-id", help="the sheet's Drive id, for the record (stays in raw.json)")
    ap.add_argument("--modified-utc", help="when the sheet was last edited, if known")
    ap.add_argument("--out", help="default: <folder>\\raw.json")
    ap.add_argument("--strict", action="store_true", help="exit with status 3 when a tab is missing or not recognised")
    a = ap.parse_args(argv)
    paths = F.listing_paths(a.folder)
    listing = F.load_json(paths["listing"], default={})
    list_date = a.list_date or listing.get("list_date")
    if not list_date:
        raise SystemExit("the list date is needed to place dates typed without a year: pass --list-date or write listing.json first")
    as_of = a.as_of or datetime.date.today().isoformat()
    slug = a.slug or listing.get("slug") or C.slugify(os.path.basename(paths["folder"]))
    sheet = listing.get("sheet") or {}
    source = {"kind": a.kind, "title": a.title or sheet.get("title"), "drive_id": a.drive_id or sheet.get("drive_id"),
              "modified_utc": a.modified_utc}
    raw = parse_workbook(a.xlsx, list_date, as_of, slug, source)
    out = a.out or paths["raw"]
    if os.path.exists(out):
        try:
            raw["since_last_parse"] = changes_since(F.load_json(out), raw)
        except Exception:                                    # an unreadable old file is replaced, not trusted
            raw["since_last_parse"] = None
    F.save_json(out, raw)
    F.say("raw.json written (PRIVATE):", out)
    F.say(summary(raw))
    for w in raw["warnings"]:
        F.shout(w)
    odd = raw["tabs_missing"] or [t for t in raw["tabs"] if t["kind"] == "other"]
    return 3 if (a.strict and odd) else 0


if __name__ == "__main__":
    sys.exit(main())
