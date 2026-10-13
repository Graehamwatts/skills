# -*- coding: utf-8 -*-
"""Builds a SYNTHETIC "View Activity" workbook that reproduces every ugly shape the real sheets have.

    python make_workbook.py <out.xlsx>

Every name, phone number, email address, brokerage and note in it is invented. The phone numbers are in the
555-01xx range, which is never assigned. Nothing here comes from a real listing, and nothing real may be added.

What is ugly in it, on purpose:
  - tab names that are not the "usual" ones (a trailing space, a suffix in brackets), an unrecognised tab,
    and a title row above the header on the first tab;
  - a Notes header merged across four columns on one tab, and unnamed spill columns with a gap on another;
  - dated lines typed as "3/13 - text", "3/16-text", "4/7 --text", "4/7this", "4/8/2026" alone on a line,
    a leading quote mark, a leading space, "3/18:text", a ">" reply line, and lines that are NOT dates
    ("2-3 buyers", "1. When"), plus a date that fits nowhere in the listing's life;
  - showing dates as true date cells, month names with ordinals, zero-width characters, "3/21/2026 at 7PM", "tbd";
  - open house rows with two sessions in one cell, a misspelled month, no date at all, a non-breaking space,
    an en dash in the time, counts in prose ("7 parties", "no one came", "only 1 group", per-day counts), and a
    numeric count column;
  - phone numbers stored as floats, a phone number typed into the email column, a typo in a date column ("4/009");
  - the four columns the team is adding (price reaction, interest level, offer intent, next follow-up date);
  - a row that only holds a row number, a row that repeats the header, "Invited by" rows, a row with no name,
    the same person on two tabs (once matched by phone, once only by name), and the same attempt typed on two tabs.

The module also gives the tests a listing.json (LISTING) and a tag file for the workbook (tags_for).
"""
import datetime
import os
import site
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.normpath(os.path.join(HERE, "..", "..", "scripts")))
if site.getusersitepackages() not in sys.path:
    sys.path.append(site.getusersitepackages())
import sr_notes as N  # noqa: E402

D = datetime.datetime
ZW = "\u200c"
AS_OF = "2026-04-13"
REPORT_DATE = "2026-04-13"          # a Monday; the report covers Mon Apr 6 to Sun Apr 12
LISTING = {
    "schema_version": 1, "slug": "sample-glen",
    "address": "218 Sample Glen Ct", "city": "Fixture City", "state": "CA", "zip": "94000",
    "beds": 3, "baths": 2.0, "sqft": 1480, "property_type": "Single Family",
    "list_date": "2026-03-12", "list_price": 1175000,
    "price_history": [{"date": "2026-03-12", "price": 1195000, "event": "listed"},
                      {"date": "2026-04-01", "price": 1175000, "event": "price_change"}],
    "status": "Active", "status_checked": "2026-04-13", "mls": "ML00000218", "hero_photo": None, "property_page": None,
    "seller_first_names": ["Pat", "Sam"],
    "numbering": {"first_issue_number": 5, "first_issue_date": "2026-04-13", "note": None},
    "sheet": {"drive_id": "synthetic-drive-id", "title": "218 Sample Glen Ct - View Activity"},
    "keybox_number": "00000000", "listing_folder": None,
    "tracking": {"keybox_showings": False, "disclosure_sent": False, "disclosure_opened": False},
}

GLIDE_HEAD = ["Agent #", "Buyer #", "Agent name", "Phone number", "Email address", "Feedback", "Email", "Call", "Text",
              "Notes", None, None, None, "Price reaction", "Interest level", "Offer intent", "Next follow-up date"]
GLIDE = [
    [1.0, None, "Odalys Verhoeven", 6505550101.0, "odalys@harborlight.example.com", "refer to showing activity for updates",
     None, D(2026, 4, 13), D(2026, 4, 7),
     "3/13 - no answer, left a vm. sent a text and email", "3/16-no answer, left a vm. sent a text and email",
     "4/7 --called about the price change, she will ask her clients", "4/13 - checked in, she said they may write an offer this week",
     "Too high", "Warm", "Maybe", D(2026, 4, 15)],
    [2.0, None, "Thaddeus Okonkwo", "(650) 555-0102", "thaddeus@example.com", "3/20 - reviewing the disclosures with his clients",
     None, None, None, "3/19 - sent the disclosure link"],
    [None, 1.0, "Invited by Thaddeus Okonkwo"],
    [None, 2.0, "Invited by Thaddeus Okonkwo"],
    [3.0, None, "Leocadia Marchetti", "650.555.0103", "leocadia@example.com", "open house host",
     None, None, None, "3/15 - no answer, left a vm. sent a text and email"],
    [6.0],
    [None, None, "Kestrel", None, 6505550104.0, "This is a buyer lead who reached out to the team\nrequested the disclosure package",
     None, None, None, "4/7this is a buyer, she called back and asked about the fence",
     "4/8/2026\n-is the fence shared with the neighbor?\n>answered by text"],
    [4.0, None, "Wilhelmina Ruiz-Castell", 6505550105.0, "wilhelmina@example.com",
     "\"4/2 - client liked the home but is not ready to make an offer\n 4/9Our clients have decided to take a break from looking",
     None, "4/009", None, "4/1 - hang up on me, sent a text and email", "2-3 buyers of mine may be interested\n1. When are offers reviewed"],
    [5.0, None, "Cosima" + ZW + " Fairweather", 6505550106.0, "cosima@example.com", None,
     None, None, None, "4/10 - no answer, left a vm. sent a text and email\n>she replied by text that her client is out of town"],
    [7.0, None, "Quentin Abernathy-Solis", 6505550112.0, "quentin@example.com", None,
     None, None, None, "4/13 - no answer, left a vm. sent a text and email"],
]
SHOW_HEAD = ["Agent #", "Buyer #", "Agent name", "Buyer Name", "Phone number", "Email address", "Date & Time", "Feedback",
             "Email", "Call", "Text", "Notes"]
SHOW = [
    [None, None, "Odalys Verhoeven", None, "650-555-0101", "odalys@harborlight.example.com", D(2026, 3, 14, 14, 0),
     "3/15 - her clients liked the layout and the yard\n3/22- still thinking it over, said the price feels high for the street",
     None, None, None, "3/16-no answer, left a vm. sent a text and email"],
    [None, None, "Bartholomew Nguyen-Hale", None, 6505550107.0, "bartholomew@example.com",
     "Tuesday, Mar 17th" + ZW + " at " + ZW + "11:35 am" + ZW, "3/18:said the price is fair for the area and the layout works",
     None, None, None, None, None, None, "3/19 - sent a follow up text"],
    [None, None, "Ignatius Delacroix", None, 6505550108.0, "ignatius@example.com", "3/21/2026 at 7PM",
     "3/22 - not interested, client said the kitchen is awful and it is way overpriced",
     None, None, None, "3/23 - no answer, left a vm. sent a text and email", "3/25 - number not in service, sent an email",
     "3/27 - number not in service, sent an email", "3/30 - sent email, number no longer in service"],
    [None, None, None, None, 6505550109.0, None, None, "4/6-wanted to show the home but the client liked another home better",
     None, D(2026, 4, 6)],
    [None, None, "Persephone Adeyemi", None, 6505550110.0, "persephone@example.com", "Apr 8th at 2:00 pm",
     "4/9 - clients want to write an offer this week", None, D(2026, 4, 9)],
    [None, None, "Thaddeus Okonkwo", None, None, None, "Mar 28th" + ZW + " at " + ZW + "1:00 pm",
     "3/29 - clients loved the yard, the stairs are a problem for them\n4/8 - they chose another home"],
    ["Agent #", "Buyer #", "Agent name", "Buyer Name", "Phone number"],
    [None, None, "Zenobia Achterberg", None, 6505550111.0, "zenobia@example.com", "tbd", "12/30 - left over from another sheet",
     None, None, None, "4/9 - no answer, left a vm. sent a text and email", "4/10 - no answer, left a vm. sent a text and email",
     "4/11 - no answer, left a vm. sent a text and email"],
]
OH_HEAD = ["Agent #", "Agent Name", "Phone number", "Email Address", "Date of Open House", "Feedback", "Visitors"]
OPEN_HOUSE = [
    [None, "Leocadia Marchetti", "650.555.0103", "leocadia@example.com", "Mar 14, Sat 12-2PM", "7 parties came to visit, asked about the HOA"],
    [None, "Leocadia Marchetti", None, None, "Sat 2-4PM"],
    [None, "Leocadia Marchetti", None, None, "3/21 and 3/22, 1-4PM",
     "Saturday: No one came\nSunday: open house was good. We had two guests. One asked for the disclosures\n-sent her the disclosure package"],
    [None, "Octavian Brightwater", 6505550113.0, None, "Marhc 28, 1:30-4:30PM", "only 1 group drove by, no feedback"],
    [None, "Leocadia Marchetti", None, None, "3/29\u00a0(Sunday) 1-4PM",
     "6 groups, one serious buyer with an agent. last group felt the price was above what the market supports"],
    [None, "Leocadia Marchetti", None, None, D(2026, 4, 4), "No one came"],
    [None, "Leocadia Marchetti", None, None, "Apr 5, Sun| 1:00 PM \u2013 4:00 PM", "quiet afternoon", 3.0],
    [None, "Leocadia Marchetti", None, None, "4/11, 1-4PM\n4/12, 11AM-1PM", "2 visitors on saturday, 1 on sunday"],
]
INQ_HEAD = ["Agent name", "Buyer Name", "Phone number", "Email address", "Feedback", "Email", "Call", "Text", "Note"]
INQUIRY = [[None, "Kestrel", 6505550104.0, None, "4/9-asked if the seller would consider a rent back"]]
OFFER_HEAD = ["Agent #", "Agent Name", "Phone number", "Email Address", "Offer Summary", "Received Date"]
OFFERS = [[None, "Persephone Adeyemi", 6505550110.0, "persephone@example.com", "Offer at 1,150,000, conventional, 20% down, 21 day close", D(2026, 4, 10)]]

TITLES = {"glide": "Glide View Activity ", "showing": "Showing Activity (Suprakey)", "open_house": "Open House Feedback",
          "inquiry": "Inquiry", "offers": "Offers Received", "other": "Team notes"}


def build(path, shuffle=False, drop=(), titles=None, tag_columns=True, edits=None):
    """Writes the workbook. shuffle reverses the data rows of every tab and puts a blank row among them, which
    must not change a single key. drop leaves tabs out. titles renames tabs. tag_columns=False writes the sheet
    as it was before the team added its four columns. edits maps (tab, row index, column index) to a new value."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    names = dict(TITLES)
    names.update(titles or {})
    edits = edits or {}

    def rows_of(key, data):
        data = [list(r) for r in data]
        for (tab, ri, ci), v in edits.items():
            if tab == key:
                data[ri] = data[ri] + [None] * (ci + 1 - len(data[ri]))
                data[ri][ci] = v
        if shuffle:
            data = list(reversed(data))
            data.insert(1, [])
        return data

    if "glide" not in drop:
        ws = wb.create_sheet(names["glide"])
        ws.append(["218 Sample Glen Ct, who has the disclosure package"])
        ws.append([])
        head = GLIDE_HEAD if tag_columns else GLIDE_HEAD[:13]
        ws.append(head)
        ws.merge_cells("J3:M3")
        for r in rows_of("glide", GLIDE):
            ws.append(r if tag_columns else r[:13])
    if "showing" not in drop:
        ws = wb.create_sheet(names["showing"])
        ws.append(SHOW_HEAD)
        for r in rows_of("showing", SHOW):
            ws.append(r)
    if "open_house" not in drop:
        ws = wb.create_sheet(names["open_house"])
        ws.append(OH_HEAD if tag_columns else OH_HEAD[:6])
        for r in rows_of("open_house", OPEN_HOUSE):
            ws.append(r if tag_columns else r[:6])
    if "inquiry" not in drop:
        ws = wb.create_sheet(names["inquiry"])
        ws.append(INQ_HEAD)
        for r in rows_of("inquiry", INQUIRY):
            ws.append(r)
    if "offers" not in drop:
        ws = wb.create_sheet(names["offers"])
        ws.append(OFFER_HEAD)
        for r in rows_of("offers", OFFERS):
            ws.append(r)
    if "other" not in drop:
        ws = wb.create_sheet(names["other"])
        ws.append(["remember to order the open house signs"])
    wb.save(path)
    return path


# ---------------------------------------------------------------------------------------------
# a tag file for the workbook, as Claude would write it after reading the todo list
# ---------------------------------------------------------------------------------------------
def IN(text, channel="unknown"):
    return {"channel": channel, "direction": "inbound", "outcome": "replied", "summary_safe": text}


def OUT(text, channel, outcome):
    return {"channel": channel, "direction": "outbound", "outcome": outcome, "summary_safe": text}


def FB(text, price="not_raised", softened=False, themes=(), liked=()):
    return {"text_safe": text, "softened": softened, "source": "conversation", "counts_as_conversation": True,
            "price_reaction": price, "price_asked": None, "themes": list(themes), "liked": list(liked)}


VM, TX, EM = dict(N.VOICEMAIL), dict(N.TEXT), dict(N.EMAIL)
# a phrase from the raw note -> how that note is tagged. Stock phrases are not here: the rules read them.
CANNED = [
    ("her clients liked the layout", [IN("Said her clients liked the layout and the yard.")],
     FB("Her clients liked the layout and the yard.", liked=["layout_space", "outdoor_space"])),
    ("price feels high", [IN("Said her clients are still thinking it over and feel the price is high for the street.")],
     FB("Her clients are still thinking it over. They feel the price is high for the street.", price="too_high", themes=["still_deciding"])),
    ("called about the price change", [OUT("Called about the price change. She will ask her clients.", "call", "reached")], None),
    ("may write an offer this week", [OUT("Checked in. She said her clients may write an offer this week.", "call", "reached")], None),
    ("reviewing the disclosures", [IN("Said he is reviewing the disclosures with his clients.")], None),
    ("sent the disclosure link", [OUT("Sent the disclosure package link.", "email", "sent")], None),
    ("loved the yard", [IN("Said his clients loved the yard. The stairs are a problem for them.")],
     FB("His clients loved the yard. The stairs are a problem for them.", themes=["stairs"], liked=["outdoor_space"])),
    ("they chose another home", [IN("Said his clients chose another property.")],
     FB("His clients chose another property.", themes=["chose_another_home"])),
    ("This is a buyer lead", [], None),
    ("asked about the fence", [IN("Called back and asked about the fence.", "call")], None),
    ("is the fence shared", [IN("Asked whether the fence is shared with the neighbor."), OUT("Answered her question by text.", "text", "sent")], None),
    ("rent back", [IN("Asked whether the seller would consider a rent back.")], None),
    ("not ready to make an offer", [IN("Said her client liked the home but is not ready to make an offer.")],
     FB("Her client liked the home but is not ready to make an offer.", themes=["still_deciding"], liked=["home_overall"])),
    ("take a break", [IN("Said her clients have decided to take a break from looking.")],
     FB("Her clients have decided to take a break from looking.", themes=["buyer_timing"])),
    ("2-3 buyers of mine", [], None),
    ("out of town", [VM, TX, EM, IN("Replied by text that her client is out of town.", "text")], None),
    ("price is fair", [IN("Said the price is fair for the area and the layout works.")],
     FB("Said the price is fair for the area and that the layout works for his clients.", price="fair_or_positive", liked=["layout_space"])),
    ("kitchen is awful", [IN("Said his client is not interested.")],
     FB("His client is not interested. He gave negative feedback on the kitchen and specifically said the price is too high.",
        price="too_high", softened=True, themes=["condition_disclosures"])),
    ("wanted to show the home", [{"channel": "call", "direction": "unknown", "outcome": "reached",
                                  "summary_safe": "Spoke by phone. The agent had wanted to show the home, but the client preferred another property."}],
     FB("Had wanted to show the home, but the client preferred another property.", themes=["chose_another_home"])),
    ("want to write an offer", [IN("Said her clients want to write an offer this week.", "call")],
     FB("Her clients want to write an offer this week.", liked=["home_overall"])),
    ("left over from another sheet", [], None),
]


def tags_for(raw, include_late_person=True):
    """The tag file for a raw.json parsed from this workbook: who is who, and how each note reads to a seller."""
    rows = [(t["kind"], r) for t in raw["tabs"] for r in t["rows"]]

    def keys(name, kinds=("disclosure", "showing", "inquiry")):
        return [r["key"] for k, r in rows if k in kinds and (r.get("name_raw") or r.get("buyer_raw") or "").replace(ZW, "") == name]

    host_rows = keys("Leocadia Marchetti") + keys("Leocadia Marchetti", ("open_house",))
    invited = [r["key"] for k, r in rows if r.get("invited_by_raw")]
    no_name = [r["key"] for k, r in rows if k == "showing" and not r.get("name_raw") and r.get("phones")]
    zen = keys("Zenobia Achterberg")
    people = {
        "c001": {"rows": keys("Odalys Verhoeven"), "kind": "agent", "brokerage": "Harborlight Realty",
                 "tags": {"interest": "cold", "offer_intent": "unknown"},
                 "feedback_safe": "Her clients liked the layout and the yard. They are still thinking it over and feel the price is high for the street."},
        "c002": {"rows": keys("Thaddeus Okonkwo"), "kind": "agent", "tags": {"interest": "passed", "offer_intent": "no"}, "status": "passed",
                 "feedback_safe": "His clients loved the yard, but the stairs are a problem for them. They chose another property."},
        "c003": {"rows": invited[:1], "kind": "buyer", "represented_by": "c002"},
        "c004": {"rows": invited[1:], "kind": "buyer", "represented_by": "c002"},
        "c005": {"rows": host_rows, "kind": "agent", "is_open_house_host": True},
        "c006": {"rows": keys("Kestrel"), "kind": "buyer", "tags": {"interest": "warm"}},
        "c007": {"rows": keys("Wilhelmina Ruiz-Castell"), "kind": "agent", "tags": {"interest": "cold", "offer_intent": "no"},
                 "feedback_safe": "Her client liked the home but was not ready to make an offer. They have since decided to take a break from looking."},
        "c008": {"rows": keys("Cosima Fairweather"), "kind": "agent"},
        "c009": {"rows": keys("Bartholomew Nguyen-Hale"), "kind": "agent", "tags": {"interest": "warm", "offer_intent": "unknown"},
                 "feedback_safe": "Said the price is fair for the area and that the layout works for his clients."},
        "c010": {"rows": keys("Ignatius Delacroix"), "kind": "agent", "tags": {"interest": "passed", "offer_intent": "no"}, "status": "passed",
                 "feedback_safe": "His client is not interested. He gave negative feedback on the kitchen and specifically said the price is too high.",
                 "feedback_softened": True},
        "c011": {"rows": no_name, "kind": "agent", "tags": {"interest": "passed", "offer_intent": "no"}, "status": "passed",
                 "feedback_safe": "Had wanted to show the home, but the client preferred another property.",
                 "engagements": dict((k, [{"type": "inquiry", "date": "2026-04-06", "date_quality": "exact"}]) for k in no_name)},
        "c012": {"rows": keys("Persephone Adeyemi"), "kind": "agent", "brokerage": "Tidewater Homes",
                 "tags": {"interest": "hot", "offer_intent": "yes"}, "status": "offer_submitted",
                 "feedback_safe": "Her clients want to write an offer this week.",
                 "offer_basis_safe": "Said her clients want to write an offer this week.", "offer_basis_date": "2026-04-09"},
        "c013": {"rows": zen, "kind": "agent",      # the date cell says "tbd": the notes show this row was an inquiry
                 "engagements": dict((k, [{"type": "inquiry", "date": "2026-04-09", "date_quality": "approximate"}]) for k in zen)},
    }
    if include_late_person:
        people["c014"] = {"rows": keys("Quentin Abernathy-Solis"), "kind": "agent"}
    entries = {}
    held = set(k for p in people.values() for k in p["rows"])
    for kind, r in rows:
        if r["key"] not in held or kind == "open_house":
            continue
        for e in r["entries"]:
            touches, formulaic = N.infer_touches(e["text"]) if e["field"] == "note" else ([], False)
            if (formulaic and e["date"]) or (not e["date"] and N.is_label(e["text"])):
                continue
            hit = [c for c in CANNED if c[0] in e["text"]]
            if not hit:
                raise AssertionError("make_workbook.CANNED has no reading for the note: %r" % e["text"][:60])
            entries[e["key"]] = {"touches": [dict(t) for t in hit[0][1]], "feedback": dict(hit[0][2]) if hit[0][2] else None}
    sessions = dict((s["key"], (r, s)) for k, r in rows if k == "open_house" for s in r["sessions"])
    open_houses = {}
    for key, (r, s) in sessions.items():
        if s["date"] is None:
            open_houses[key] = {"date": "2026-03-14", "date_quality": "approximate",
                                "date_note": "No date is typed. Placed on Mar 14 by its position under the Mar 14 row."}
        elif s["date"] == "2026-03-29":
            open_houses[key] = {"recap_safe": "6 groups. One was a serious buyer with an agent. The last group felt the price was above what the market supports.",
                                "softened": True, "price_remark": {"reaction": "too_high", "softened": True,
                                                                   "remark_safe": "One open house group felt the price was above what the market supports."}}
        elif s["date"] == "2026-03-22":
            open_houses[key] = {"recap_safe": "Two guests. One asked for the disclosure package.", "disclosure_requests": 1}
    offer_rows = [r["key"] for k, r in rows if k == "offers"]
    offers = [{"id": "o1", "date": "2026-04-10", "contact_id": "c012", "price": 1150000, "state": "received", "source_row": offer_rows[0],
               "terms_summary": "Conventional loan, 20% down, 21 day close."}] if offer_rows else []
    return {"people": people, "entries": entries, "open_houses": open_houses, "offers": offers,
            "next_week": ["Check back with Odalys V. on her clients' decision."]}


def prepare(folder, tagged=True, through=None, **build_kw):
    """A complete synthetic listing folder for the tests: the workbook, listing.json, raw.json and (when tagged)
    curation.json. Returns (raw, curation or None). The folder must be outside every git repo (use a temp folder)."""
    import intake_sheet
    import curation_tools
    import sr_files as F
    os.makedirs(folder, exist_ok=True)
    xlsx = build(os.path.join(folder, "export.xlsx"), **build_kw)
    F.save_json(os.path.join(folder, "listing.json"), LISTING)
    raw = intake_sheet.parse_workbook(xlsx, LISTING["list_date"], AS_OF, LISTING["slug"], {"title": LISTING["sheet"]["title"]})
    F.save_json(os.path.join(folder, "raw.json"), raw)
    if not tagged:
        return raw, None
    cur = curation_tools.blank_curation(LISTING["slug"])
    report = curation_tools.apply_tags(cur, tags_for(raw), raw, today=AS_OF)
    if report["rejected"]:
        raise AssertionError("the synthetic readings file was rejected: %s" % report["rejected"][:3])
    F.save_json(os.path.join(folder, "curation.json"), cur)
    return raw, cur


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.getcwd(), "sample-glen-view-activity.xlsx")
    print("wrote", build(out))
