# -*- coding: utf-8 -*-
"""Tests for intake_sheet.py and sr_sheet_text.py: reading a hand-kept View Activity workbook.

    python -m unittest tests.test_intake        (from the skill folder)
    python tests\\test_intake.py

Everything here is synthetic. The workbook comes from tests/fixtures/make_workbook.py and is written to a
temp folder, never into the repo.
"""
import contextlib
import datetime
import io
import json
import os
import shutil
import site
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "scripts"))
sys.path.insert(0, os.path.join(HERE, "fixtures"))
if site.getusersitepackages() not in sys.path:
    sys.path.append(site.getusersitepackages())
import sr_contract as C  # noqa: E402
import sr_files as F  # noqa: E402
import sr_sheet_text as T  # noqa: E402
import intake_sheet  # noqa: E402
import make_workbook as W  # noqa: E402

try:
    import openpyxl
except ImportError:                                     # the parser needs it; without it these tests cannot run
    openpyxl = None

CTX = T.DateCtx("2026-03-12", "2026-04-13")


@contextlib.contextmanager
def quiet():
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        yield


def run(argv):
    with quiet():
        return intake_sheet.main(argv)

DASH = "\u2013"


def dated(text, ctx=CTX):
    return [(e["date"], e["text"]) for e in T.split_dated(text, ctx)]


class Cells(unittest.TestCase):
    def test_zero_width_characters_and_odd_spaces_are_removed(self):
        self.assertEqual(T.clean("Mar 17th\u200c at \u200c11:35 am\u200c"), "Mar 17th at 11:35 am")
        self.assertEqual(T.clean("3/29\u00a0(Sunday)\r\n"), "3/29 (Sunday)")
        self.assertEqual(T.clean("\ufeff  "), None)
        self.assertEqual(T.clean("client\u2019s"), "client's")

    def test_a_phone_number_stored_as_a_float_reads_as_digits(self):
        self.assertEqual(T.clean(6505550101.0), "6505550101")
        self.assertEqual(T.phones_of(6505550101.0, "(650) 555-0102 or 650.555.0103", "+1 650 555 0104"),
                         ["6505550101", "6505550102", "6505550103", "6505550104"])
        self.assertEqual(T.phones_of("agent6505550101@example.com"), [], "digits inside an email address are not a phone")

    def test_dates_pass_through_and_serialise(self):
        d = datetime.datetime(2026, 3, 14, 14, 0)
        self.assertIs(T.clean(d), d)
        self.assertEqual(T.cell_json(d), "2026-03-14T14:00:00")

    def test_name_key_is_first_and_last(self):
        self.assertEqual(T.name_key("Robert (Bob) Hale\nsecond line"), "robert_hale")
        self.assertEqual(T.name_key("Kestrel"), "kestrel")
        self.assertIsNone(T.name_key("650-555-0101"))


class DatedLines(unittest.TestCase):
    def test_every_shape_the_team_types(self):
        cases = {
            "3/13 - no answer": ("2026-03-13", "no answer"),
            "3/16-no answer": ("2026-03-16", "no answer"),
            "4/7 --called her": ("2026-04-07", "called her"),
            "4/7this is a buyer": ("2026-04-07", "this is a buyer"),
            "\"4/2 - client liked it": ("2026-04-02", "client liked it"),
            " 4/9Our clients decided": ("2026-04-09", "Our clients decided"),
            "3/18:said the price is fair": ("2026-03-18", "said the price is fair"),
            "4/8 not a good time": ("2026-04-08", "not a good time"),
            "4/009 no answer": ("2026-04-09", "no answer"),
            "- 4/3 - texted": ("2026-04-03", "texted"),
            "4/8/2026 - with a year": ("2026-04-08", "with a year"),
            "4/8/26 - short year": ("2026-04-08", "short year"),
            "3/20 " + DASH + " a typed dash": ("2026-03-20", "a typed dash"),
        }
        for text, want in cases.items():
            self.assertEqual(dated(text), [want], text)

    def test_a_date_alone_on_a_line_takes_the_lines_under_it(self):
        got = T.split_dated("4/8/2026\n-is the fence shared?\n>answered by text", CTX)
        self.assertEqual(len(got), 1)
        self.assertEqual(got[0]["date"], "2026-04-08")
        self.assertEqual(got[0]["year_rule"], "year_typed")
        self.assertEqual(got[0]["text"], "-is the fence shared?\n>answered by text")
        self.assertEqual(got[0]["reply_lines"], ["answered by text"])

    def test_lines_that_only_look_like_dates(self):
        for text in ("2-3 buyers of mine may be interested", "1. When are offers reviewed", "1- Any pending litigation",
                     "24/7 access to the gym", "1/2 bath downstairs", "13/45 is not a date", "466 Example Ave"):
            self.assertEqual(dated(text), [(None, text)], text)

    def test_undated_lead_in_and_running_cell(self):
        got = dated("refer to showing activity for updates\n3/20\nclient liked it a lot\n3/29 - they passed")
        self.assertEqual(got, [(None, "refer to showing activity for updates"), ("2026-03-20", "client liked it a lot"),
                               ("2026-03-29", "they passed")])

    def test_the_year_comes_from_the_listings_life_not_from_today(self):
        ctx = T.DateCtx("2025-11-14", "2026-03-05")              # a listing that runs across New Year
        self.assertEqual(dated("12/30 - showed it", ctx), [("2025-12-30", "showed it")])
        self.assertEqual(dated("1/5 - followed up", ctx), [("2026-01-05", "followed up")])
        self.assertEqual(dated("11/1 - a call before it was listed", ctx), [("2025-11-01", "a call before it was listed")])

    def test_a_date_outside_the_listings_life_is_flagged_not_trusted(self):
        e = T.split_dated("12/30 - left over from another sheet", CTX)[0]
        self.assertEqual((e["date"], e["year_rule"]), ("2025-12-30", "outside_listing_life"))

    def test_a_second_date_inside_a_line_is_reported_not_split(self):
        e = T.split_dated("3/21 - sent the link 3/28-client is not interested", CTX)
        self.assertEqual(len(e), 1)
        self.assertEqual(e[0]["mid_line_dates"], ["3/28"])


class ShowingDates(unittest.TestCase):
    def when(self, v):
        return [(w["date"], w["time_text"], w["date_quality"]) for w in T.parse_when(v, CTX)]

    def test_formats(self):
        self.assertEqual(self.when(datetime.datetime(2026, 3, 14, 14, 0)), [("2026-03-14", "2 PM", "exact")])
        self.assertEqual(self.when(datetime.datetime(2026, 3, 14)), [("2026-03-14", None, "exact")])
        self.assertEqual(self.when("Tuesday, Mar 17th\u200c at \u200c11:35 am\u200c"), [("2026-03-17", "11:35 AM", "exact")])
        self.assertEqual(self.when("3/21/2026 at 7PM"), [("2026-03-21", "7 PM", "exact")])
        self.assertEqual(self.when("Wednesday, April 1 at 2PM"), [("2026-04-01", "2 PM", "exact")])
        self.assertEqual(self.when("Sun, Mar 29th from 1:00 pm to 1:30 pm"), [("2026-03-29", "1 to 1:30 PM", "exact")])
        self.assertEqual(self.when("Sept 17, 5-6PM")[0][1], "5 to 6 PM")
        self.assertEqual(self.when("tbd"), [(None, None, "missing")])
        self.assertEqual(self.when(None), [])

    def test_a_misspelled_month_is_read(self):
        self.assertEqual(self.when("Marhc 28, 1:30-4:30PM"), [("2026-03-28", "1:30 to 4:30 PM", "exact")])
        self.assertEqual(T.month_of("Septmber"), 9)
        self.assertEqual(T.month_of("Febuary"), 2)
        self.assertIsNone(T.month_of("Sat"))
        self.assertIsNone(T.month_of("Marsh"))

    def test_two_sessions_in_one_cell(self):
        self.assertEqual(self.when("3/21 and 3/22, 1-4PM"), [("2026-03-21", "1 to 4 PM", "exact"), ("2026-03-22", "1 to 4 PM", "exact")])
        self.assertEqual(self.when("4/11, 1-4PM\n4/12, 11AM-1PM"), [("2026-04-11", "1 to 4 PM", "exact"), ("2026-04-12", "11 AM to 1 PM", "exact")])

    def test_no_date_keeps_the_weekday_and_time(self):
        w = T.parse_when("Sat 2-4PM", CTX)[0]
        self.assertEqual((w["date"], w["date_quality"], w["weekday"], w["time_text"]), (None, "missing", 5, "2 to 4 PM"))

    def test_time_ranges_never_carry_a_dash(self):
        for text, want in (("1:00 PM " + DASH + " 4:00 PM", "1 to 4 PM"), ("10-12PM", "10 AM to 12 PM"), ("12-2:30PM", "12 to 2:30 PM")):
            self.assertEqual(T.time_text(text), want)
            self.assertFalse(C.has_dash(T.time_text(text)))

    def test_date_columns_take_typed_text_too(self):
        self.assertEqual(T.parse_last_touch("4/009", CTX), (["2026-04-09"], None))
        self.assertEqual(T.parse_last_touch("3/13, 3/15 LVM", CTX)[0], ["2026-03-13", "2026-03-15"])
        self.assertEqual(T.parse_last_touch(datetime.datetime(2026, 4, 7), CTX), (["2026-04-07"], None))
        self.assertTrue(T.parse_last_touch("soon", CTX)[1])


class OpenHouseProse(unittest.TestCase):
    def test_counts_keep_the_hosts_own_unit(self):
        self.assertEqual(T.count_in_prose("7 parties came to visit")[:2], (7, "parties"))
        self.assertEqual(T.count_in_prose("6 groups, one serious buyer with an agent")[:2], (6, "groups"))
        self.assertEqual(T.count_in_prose("Only one group came in towards the end")[:2], (1, "groups"))
        self.assertEqual(T.count_in_prose("had 3 visitors. Last group wants disclosures")[:2], (3, "visitors"))
        self.assertEqual(T.count_in_prose("We had two guests")[:2], (2, "guests"))
        self.assertEqual(T.count_in_prose("slow, 4 people came through")[:2], (4, "people"))

    def test_nobody_came_is_zero_and_no_count_is_none(self):
        for text in ("no one came", "No one came to the open house", "no visitors", "nobody showed up"):
            self.assertEqual(T.count_in_prose(text)[0], 0, text)
        for text in ("pulled the disclosures for the buyers who walked in", "quiet afternoon", "no feedback from anyone",
                     "no one wanted to make an offer"):
            self.assertIsNone(T.count_in_prose(text)[0], text)

    def test_per_day_counts_go_to_the_right_session(self):
        s, problems = T.parse_sessions("3/21 and 3/22, 1-4PM", "Saturday: No one came\nSunday: We had two guests\n-sent her the package", CTX)
        self.assertEqual([(x["date"], x["groups"], x["counted_unit"]) for x in s], [("2026-03-21", 0, "visitors"), ("2026-03-22", 2, "guests")])
        self.assertEqual(s[0]["team_note"], "sent her the package")
        self.assertEqual(problems, [])
        s, _ = T.parse_sessions("4/11, 1-4PM\n4/12, 11AM-1PM", "2 visitors on saturday, 1 on sunday", CTX)
        self.assertEqual([(x["groups"], x["counted_unit"]) for x in s], [(2, "visitors"), (1, "visitors")])
        s, _ = T.parse_sessions("Saturday: 12-2PM\nSunday, 1:30-4:30PM", "one person came through Saturday and 2 on Sunday", CTX)
        self.assertEqual([(x["date"], x["groups"]) for x in s], [(None, 1), (None, 2)])

    def test_one_count_for_two_sessions_is_not_guessed(self):
        s, problems = T.parse_sessions("3/21 and 3/22", "5 groups in total", CTX)
        self.assertEqual([x["groups"] for x in s], [None, None])
        self.assertTrue(any("could not be split" in p for p in problems))

    def test_a_count_column_beats_the_prose(self):
        s, _ = T.parse_sessions("Apr 5", "quiet afternoon", CTX, count_value=3.0, count_unit="visitors")
        self.assertEqual((s[0]["groups"], s[0]["counted_unit"], s[0]["groups_source"]), (3, "visitors", "sheet_column"))
        s, _ = T.parse_sessions("Apr 5", "6 groups", CTX, count_value=4.0)
        self.assertEqual((s[0]["groups"], s[0]["counted_unit"], s[0]["groups_source"]), (4, "groups", "sheet_column"))


class TagColumns(unittest.TestCase):
    def test_the_four_columns_are_normalised(self):
        self.assertEqual(T.read_sheet_tag("price_reaction", "Too high", CTX), ("too_high", None))
        self.assertEqual(T.read_sheet_tag("price_reaction", "not discussed", CTX), ("not_raised", None))
        self.assertEqual(T.read_sheet_tag("price_reaction", "Fair", CTX), ("fair_or_positive", None))
        self.assertEqual(T.read_sheet_tag("interest", "WARM", CTX), ("warm", None))
        self.assertEqual(T.read_sheet_tag("interest", "not interested", CTX), ("passed", None))
        self.assertEqual(T.read_sheet_tag("offer_intent", "Maybe", CTX), ("maybe", None))
        self.assertEqual(T.read_sheet_tag("next_followup_date", datetime.datetime(2026, 4, 15), CTX), ("2026-04-15", None))
        self.assertEqual(T.read_sheet_tag("next_followup_date", "4/15", CTX), ("2026-04-15", None))

    def test_a_value_it_does_not_understand_is_left_for_a_person(self):
        value, problem = T.read_sheet_tag("interest", "7", CTX)
        self.assertIsNone(value)
        self.assertIn("not understood", problem)
        self.assertEqual(T.read_sheet_tag("price_reaction", None, CTX), (None, None))


@unittest.skipUnless(openpyxl, "openpyxl is not installed")
class Workbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="sr_intake_")
        cls.xlsx = W.build(os.path.join(cls.tmp, "export.xlsx"))
        cls.raw = intake_sheet.parse_workbook(cls.xlsx, W.LISTING["list_date"], W.AS_OF, "sample-glen")
        cls.tabs = dict((t["kind"], t) for t in cls.raw["tabs"])

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def row(self, kind, name):
        return next(r for r in self.tabs[kind]["rows"] if (r["name_raw"] or r["buyer_raw"] or "") == name)

    def parse(self, **kw):
        path = W.build(os.path.join(self.tmp, "variant_%d.xlsx" % len(os.listdir(self.tmp))), **kw)
        return intake_sheet.parse_workbook(path, W.LISTING["list_date"], W.AS_OF, "sample-glen")

    def test_tabs_are_found_by_keyword_not_by_exact_name(self):
        self.assertEqual([(t["title"], t["kind"]) for t in self.raw["tabs"]],
                         [("Glide View Activity ", "disclosure"), ("Showing Activity (Suprakey)", "showing"), ("Open House Feedback", "open_house"),
                          ("Inquiry", "inquiry"), ("Offers Received", "offers"), ("Team notes", "other")])
        renamed = self.parse(titles={"glide": "Disclosure pulls", "showing": "SUPRA showings", "open_house": "OH recap", "inquiry": "Coming-soon inquiries"})
        self.assertEqual([t["kind"] for t in renamed["tabs"]][:4], ["disclosure", "showing", "open_house", "inquiry"])
        self.assertEqual(renamed["tabs_missing"], [])
        self.assertEqual(sorted(r["key"] for t in renamed["tabs"] for r in t["rows"]), sorted(r["key"] for t in self.raw["tabs"] for r in t["rows"]),
                         "a renamed tab changes no key")

    def test_a_missing_or_unrecognised_tab_is_reported_loudly(self):
        raw = self.parse(drop=("showing",))
        self.assertEqual(raw["tabs_missing"], ["showing"])
        self.assertTrue(any("no showing tab was found" in w for w in raw["warnings"]))
        self.assertTrue(any("Team notes" in w for w in self.raw["warnings"]))
        self.assertEqual(self.tabs["other"]["rows"], [])
        self.assertTrue(any(s["tab"] == "Team notes" and "no header row" in s["reason"] for s in self.raw["skipped_rows"]))

    def test_the_header_row_is_found_under_a_title_row(self):
        t = self.tabs["disclosure"]
        self.assertEqual(t["header_row"], 3)
        self.assertEqual(t["columns_found"]["name"], "C")
        self.assertEqual(t["columns_found"]["feedback"], ["F"])
        self.assertEqual(len(t["rows"]), 9)
        self.assertTrue(any(s["row"] == 1 and "above the header row" in s["reason"] for s in self.raw["skipped_rows"]))

    def test_every_notes_column_is_read_to_the_last_used_column(self):
        self.assertEqual(self.tabs["disclosure"]["columns_found"]["notes"], ["J", "K", "L", "M"], "a Notes header merged across four columns")
        self.assertEqual(self.tabs["showing"]["columns_found"]["notes"], ["L", "M", "N", "O"], "unnamed spill columns to the right of Notes")
        odalys = self.row("disclosure", "Odalys Verhoeven")
        self.assertEqual([(e["col"], e["date"]) for e in odalys["entries"] if e["field"] == "note"],
                         [("J", "2026-03-13"), ("K", "2026-03-16"), ("L", "2026-04-07"), ("M", "2026-04-13")])
        bart = self.row("showing", "Bartholomew Nguyen-Hale")
        self.assertEqual([(e["col"], e["date"], e["text"]) for e in bart["entries"] if e["field"] == "note"],
                         [("O", "2026-03-19", "sent a follow up text")], "a note three columns past the header, with empty columns between")
        ignatius = self.row("showing", "Ignatius Delacroix")
        self.assertEqual(len([e for e in ignatius["entries"] if e["field"] == "note"]), 4)

    def test_no_row_is_dropped_silently(self):
        wb = openpyxl.load_workbook(self.xlsx, data_only=True)
        for ws in wb.worksheets:
            filled = [i for i, row in enumerate(ws.iter_rows(values_only=True), 1) if any(v is not None for v in row)]
            tab = next(t for t in self.raw["tabs"] if t["title"] == ws.title)
            read = [r["row"] for r in tab["rows"]]
            skipped = [s["row"] for s in self.raw["skipped_rows"] if s["tab"] == ws.title]
            header = [tab["header_row"]] if tab["header_row"] else []
            self.assertEqual(sorted(read + skipped + header), filled, ws.title)
        reasons = sorted(s["reason"] for s in self.raw["skipped_rows"])
        self.assertEqual(len(reasons), 4)
        self.assertTrue(any("only a row number" in r for r in reasons))
        self.assertTrue(any("repeats the header row" in r for r in reasons))
        self.assertTrue(all(s["cells"] for s in self.raw["skipped_rows"]), "a skipped row keeps its cells")

    def test_dated_lines_in_the_workbook(self):
        w = self.row("disclosure", "Wilhelmina Ruiz-Castell")
        self.assertEqual([(e["field"], e["date"]) for e in w["entries"]],
                         [("feedback", "2026-04-02"), ("feedback", "2026-04-09"), ("note", "2026-04-01"), ("note", None)])
        self.assertEqual(w["entries"][3]["text"], "2-3 buyers of mine may be interested\n1. When are offers reviewed")
        self.assertEqual(w["last_touch"]["call"], "2026-04-09", "the typo 4/009 in the Call column")
        k = self.row("disclosure", "Kestrel")
        self.assertEqual([(e["date"], e["year_rule"]) for e in k["entries"]],
                         [(None, None), ("2026-04-07", "listing_life"), ("2026-04-08", "year_typed")])
        self.assertEqual(k["entries"][2]["reply_lines"], ["answered by text"])
        z = self.row("showing", "Zenobia Achterberg")
        self.assertEqual(z["entries"][0]["year_rule"], "outside_listing_life")
        self.assertTrue(any("outside the listing's life" in x for x in self.raw["warnings"]))
        self.assertEqual(z["when"][0]["date"], None)
        self.assertTrue(any("holds no date that could be read" in x for x in self.raw["warnings"]))

    def test_names_phones_and_showing_dates_are_cleaned(self):
        cosima = self.row("disclosure", "Cosima Fairweather")
        self.assertEqual(cosima["phones"], ["6505550106"])
        self.assertEqual(self.row("disclosure", "Odalys Verhoeven")["phones"], ["6505550101"])
        self.assertEqual(self.row("showing", "Odalys Verhoeven")["phones"], ["6505550101"])
        self.assertEqual(self.row("disclosure", "Kestrel")["phones"], ["6505550104"], "a phone number typed into the email column")
        self.assertEqual(self.row("showing", "Bartholomew Nguyen-Hale")["when"][0]["date"], "2026-03-17")
        self.assertEqual(self.row("showing", "Bartholomew Nguyen-Hale")["when"][0]["time_text"], "11:35 AM")
        self.assertEqual(self.row("showing", "Odalys Verhoeven")["when"][0]["time_text"], "2 PM")
        blank = next(r for r in self.tabs["showing"]["rows"] if "no_name" in r["hints"])
        self.assertEqual((blank["name_raw"], blank["phones"], blank["when"]), (None, ["6505550109"], []))

    def test_open_house_rows(self):
        sessions = [s for r in self.tabs["open_house"]["rows"] for s in r["sessions"]]
        got = [(s["date"], s["groups"], s["counted_unit"], s["groups_source"]) for s in sessions]
        self.assertEqual(got, [
            ("2026-03-14", 7, "parties", "parsed_from_prose"), (None, None, None, "not_recorded"),
            ("2026-03-21", 0, "visitors", "parsed_from_prose"), ("2026-03-22", 2, "guests", "parsed_from_prose"),
            ("2026-03-28", 1, "groups", "parsed_from_prose"), ("2026-03-29", 6, "groups", "parsed_from_prose"),
            ("2026-04-04", 0, "visitors", "parsed_from_prose"), ("2026-04-05", 3, "visitors", "sheet_column"),
            ("2026-04-11", 2, "visitors", "parsed_from_prose"), ("2026-04-12", 1, "visitors", "parsed_from_prose")])
        self.assertEqual(sessions[7]["time_text"], "1 to 4 PM", "an en dash in the time cell")
        self.assertEqual(len(set(s["key"] for s in sessions)), 10)
        self.assertTrue(any("no date typed" in w for w in self.raw["warnings"]))
        self.assertEqual(self.tabs["open_house"]["rows"][0]["entries"], [], "a recap is read into sessions, not into follow-up notes")

    def test_the_teams_new_columns_are_read_when_present(self):
        o = self.row("disclosure", "Odalys Verhoeven")
        self.assertEqual(o["sheet_tags"], {"price_reaction": "too_high", "interest": "warm", "offer_intent": "maybe", "next_followup_date": "2026-04-15"})
        self.assertEqual(self.row("disclosure", "Thaddeus Okonkwo")["sheet_tags"],
                         {"price_reaction": None, "interest": None, "offer_intent": None, "next_followup_date": None})
        before = self.parse(tag_columns=False)
        rows = [r for t in before["tabs"] for r in t["rows"]]
        self.assertTrue(all(v is None for r in rows for v in r["sheet_tags"].values()))
        self.assertNotIn("price_reaction", before["tabs"][0]["columns_found"])
        self.assertEqual([s["groups"] for r in before["tabs"][2]["rows"] for s in r["sessions"]][7], None, "no count column, and the prose gives none")
        self.assertEqual(sorted(r["key"] for r in rows), sorted(r["key"] for t in self.raw["tabs"] for r in t["rows"]),
                         "adding the columns changes no key")

    def test_merges_are_proposed_never_applied(self):
        keys = [r["key"] for t in self.raw["tabs"] for r in t["rows"]]
        self.assertEqual(len(keys), len(set(keys)), "every row keeps its own key")
        self.assertEqual(self.raw["stats"]["rows_total"], 26)
        cands = dict((tuple(sorted(g["tabs"])), g) for g in self.raw["merge_candidates"])
        o1, o2 = self.row("disclosure", "Odalys Verhoeven")["key"], self.row("showing", "Odalys Verhoeven")["key"]
        t1, t2 = self.row("disclosure", "Thaddeus Okonkwo")["key"], self.row("showing", "Thaddeus Okonkwo")["key"]
        by_rows = dict((frozenset(g["rows"]), g["basis"]) for g in self.raw["merge_candidates"])
        self.assertEqual(by_rows[frozenset((o1, o2))], "phone", "the same phone as a float and as text with hyphens")
        self.assertEqual(by_rows[frozenset((t1, t2))], "name", "no phone on the second row: matched by name only")
        self.assertEqual(cands[("disclosure", "inquiry")]["basis"], "phone")
        self.assertEqual(len(self.raw["merge_candidates"]), 4)

    def test_invited_by_rows_are_kept_and_point_at_their_agent(self):
        invited = [r for r in self.tabs["disclosure"]["rows"] if r["invited_by_raw"]]
        self.assertEqual([r["invited_by_raw"] for r in invited], ["Thaddeus Okonkwo", "Thaddeus Okonkwo"])
        self.assertEqual(len(set(r["key"] for r in invited)), 2, "two buyers of one agent get two keys")
        agent = self.row("disclosure", "Thaddeus Okonkwo")["key"]
        for r in invited:
            self.assertIn(agent, self.raw["inviter_candidates"][r["key"]])

    def test_keys_survive_sorting_and_inserted_rows(self):
        shuffled = self.parse(shuffle=True)

        def keys(raw):
            return sorted(k for t in raw["tabs"] for r in t["rows"] for k in [r["key"]] + [e["key"] for e in r["entries"]] + [s["key"] for s in r.get("sessions") or []])
        self.assertEqual(keys(shuffled), keys(self.raw))
        self.assertNotEqual([r["row"] for r in shuffled["tabs"][0]["rows"]], [r["row"] for r in self.raw["tabs"][0]["rows"]])

    def test_an_edited_note_gets_a_new_key_and_nothing_else_moves(self):
        edited = self.parse(edits={("showing", 1, 7): "3/18:said the price is fair for the area, and asked about the roof"})
        old = set(e["key"] for t in self.raw["tabs"] for r in t["rows"] for e in r["entries"])
        new = set(e["key"] for t in edited["tabs"] for r in t["rows"] for e in r["entries"])
        self.assertEqual((len(old - new), len(new - old)), (1, 1))
        self.assertEqual(sorted(r["key"] for t in edited["tabs"] for r in t["rows"]), sorted(r["key"] for t in self.raw["tabs"] for r in t["rows"]))

    def test_header_variants_seen_on_other_listings(self):
        # a Notes column with unnamed spill beside it and another header further right, a showing date under another
        # name, a follow-up column the parser does not know, and no Email / Call / Text columns at all
        path = os.path.join(self.tmp, "variants.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Showing Activity (Suprashowing)"
        ws.append(["Agent #", "Agent Name", "Phone", "Email Address", "Date of Viewing/Showing", "Feedback", "Notes", None, "Lockbox # 00000000", None])
        ws.append([1.0, "Ottoline Fenwick", 6505550120.0, None, "Mar 20@ 10AM", "3/21 - liked it", "3/22 - no answer, left a vm. sent a text and email",
                   "3/24 - no answer, left a vm. sent a text and email", "on the gate", "3/26 - sent a follow up text"])
        ws = wb.create_sheet("Glide View")
        ws.append(["Agent #", "Buyer #", "Agent name", "Phone number", "Email address", "Feedback", "Follow Up Agent via"])
        ws.append([None, None, "Ottoline Fenwick", "650 555 0120", None, "3/19 - downloaded the package", "text"])
        wb.save(path)
        raw = intake_sheet.parse_workbook(path, "2026-03-12", "2026-04-13", "sample-glen", {"kind": "activity_sheet_file"})
        show, glide = raw["tabs"]
        self.assertEqual((show["kind"], glide["kind"], raw["tabs_missing"]), ("showing", "disclosure", ["open_house"]))
        self.assertEqual(show["columns_found"]["when"], "E")
        self.assertEqual(show["columns_found"]["notes"], ["G", "H", "J"])
        self.assertEqual(show["columns_unrecognised"], {"I": "Lockbox # 00000000"})
        row = show["rows"][0]
        self.assertEqual((row["when"][0]["date"], row["when"][0]["time_text"]), ("2026-03-20", "10 AM"))
        self.assertEqual([e["date"] for e in row["entries"] if e["field"] == "note"], ["2026-03-22", "2026-03-24", "2026-03-26"])
        self.assertEqual(row["cells"]["I"], "on the gate", "a cell under a header the parser does not know is kept, not interpreted")
        self.assertEqual(glide["columns_unrecognised"], {"G": "Follow Up Agent via"})
        self.assertEqual(glide["rows"][0]["last_touch"], {"email": None, "call": None, "text": None})
        self.assertTrue(any("is not a column this parser knows" in w for w in raw["warnings"]))
        self.assertEqual(raw["sources"][0]["kind"], "activity_sheet_file")
        self.assertEqual(len(raw["merge_candidates"]), 1)

    def test_what_the_old_parser_got_wrong(self):
        # it looked tabs up by exact name, wanted the header in row 1, read one Notes column, skipped "Invited by"
        # rows, dated '2-3 buyers' as Feb 3 and gave every date typed without a year the report's year
        st = self.raw["stats"]
        self.assertEqual(st["rows_by_tab_kind"]["showing"], 7)
        self.assertEqual(st["rows_by_tab_kind"]["disclosure"], 9)
        self.assertEqual(st["invited_by_rows"], 2)
        self.assertEqual(st["note_entries"], 21)
        self.assertFalse(any(e["date"] == "2026-02-03" for t in self.raw["tabs"] for r in t["rows"] for e in r["entries"]))
        self.assertEqual(st["entries_total"], st["feedback_entries"] + st["note_entries"])

    def test_raw_is_marked_private_and_carries_its_source(self):
        self.assertIn("PRIVATE", self.raw["WARNING"])
        src = self.raw["sources"][0]
        self.assertEqual((src["kind"], src["as_of"], src["file_name"]), ("activity_sheet_live", W.AS_OF, "export.xlsx"))
        self.assertEqual(len(src["sha256"]), 64)
        json.dumps(self.raw)


@unittest.skipUnless(openpyxl, "openpyxl is not installed")
class CommandLine(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="sr_intake_cli_")
        self.xlsx = W.build(os.path.join(self.tmp, "export.xlsx"))
        self.folder = os.path.join(self.tmp, "sample-glen")
        F.save_json(os.path.join(self.folder, "listing.json"), W.LISTING)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_writes_raw_json_into_the_listing_folder(self):
        code = run([self.xlsx, "--folder", self.folder, "--as-of", W.AS_OF])
        self.assertEqual(code, 0)
        raw = F.load_json(os.path.join(self.folder, "raw.json"))
        self.assertEqual((raw["slug"], raw["list_date"], raw["as_of"]), ("sample-glen", "2026-03-12", W.AS_OF))
        self.assertEqual(raw["sources"][0]["title"], "218 Sample Glen Ct - View Activity")
        self.assertEqual(run([self.xlsx, "--folder", self.folder, "--as-of", W.AS_OF]), 0)
        again = F.load_json(os.path.join(self.folder, "raw.json"))
        self.assertEqual((again["since_last_parse"]["rows_new"], again["since_last_parse"]["entries_new_or_edited"]), (0, 0))

    def test_strict_mode_fails_on_an_unrecognised_tab(self):
        self.assertEqual(run([self.xlsx, "--folder", self.folder, "--as-of", W.AS_OF, "--strict"]), 3)

    def test_private_files_are_refused_inside_a_git_repo(self):
        repo = os.path.join(self.tmp, "some-repo")
        os.makedirs(os.path.join(repo, ".git"))
        with self.assertRaises(SystemExit):
            F.save_json(os.path.join(repo, "seller-reports", "x", "raw.json"), {"a": 1})
        self.assertFalse(os.path.exists(os.path.join(repo, "seller-reports", "x", "raw.json")))
        if F.git_root(SKILL):                                   # the skill folder itself publishes to GitHub
            with self.assertRaises(SystemExit):
                F.refuse_inside_repo(os.path.join(SKILL, "tests", "fixtures", "raw.json"))

    def test_the_list_date_is_required(self):
        with self.assertRaises(SystemExit):
            run([self.xlsx, "--folder", os.path.join(self.tmp, "empty"), "--as-of", W.AS_OF])


if __name__ == "__main__":
    unittest.main()
